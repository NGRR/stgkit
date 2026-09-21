from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import shutil
import tempfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "formatos/xlsx/STGND_Datos_Graficos_Maestro_v9.xlsx"
PPTX = ROOT / "formatos/pptx/STGND_Presentacion_Maestro_v9.pptx"
LOGO = ROOT / "assets/logo-stgnd.png"

PURPLE = "4C2B46"
ORANGE = "E97700"
GRAPHITE = "222222"
LIGHT = "EFEDEF"
PETROL = "2D6673"
SAGE = "809477"
SAND = "C9B58B"
WHITE = "FFFFFF"
MIDGREY = "6F6F6F"

THEME = {
    "dk1": "000000",
    "lt1": WHITE,
    "dk2": GRAPHITE,
    "lt2": LIGHT,
    "accent1": PURPLE,
    "accent2": ORANGE,
    "accent3": PETROL,
    "accent4": SAGE,
    "accent5": SAND,
    "accent6": MIDGREY,
    "hlink": PETROL,
    "folHlink": PURPLE,
}

def patch_theme(path: Path, member: str) -> None:
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    ET.register_namespace("a", ns["a"])
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        with ZipFile(path, "r") as zin:
            zin.extractall(td)
        theme_path = td / member
        tree = ET.parse(theme_path)
        root = tree.getroot()
        scheme = root.find(".//a:clrScheme", ns)
        if scheme is not None:
            for tag, value in THEME.items():
                node = scheme.find(f"a:{tag}", ns)
                if node is None:
                    continue
                for child in list(node):
                    node.remove(child)
                srgb = ET.SubElement(node, f"{{{ns['a']}}}srgbClr")
                srgb.set("val", value)
        font_scheme = root.find(".//a:fontScheme", ns)
        if font_scheme is not None:
            for major_minor in ("majorFont", "minorFont"):
                latin = font_scheme.find(f"a:{major_minor}/a:latin", ns)
                if latin is not None:
                    latin.set("typeface", "Aptos")
        tree.write(theme_path, encoding="utf-8", xml_declaration=True)
        tmpzip = path.with_suffix(path.suffix + ".tmp")
        with ZipFile(tmpzip, "w", ZIP_DEFLATED) as zout:
            for p in td.rglob("*"):
                if p.is_file():
                    zout.write(p, p.relative_to(td).as_posix())
        tmpzip.replace(path)

def build_xlsx() -> None:
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    palette = wb.active
    palette.title = "Paleta STG"
    data = wb.create_sheet("Datos")
    dash = wb.create_sheet("Dashboard")

    palette.sheet_view.showGridLines = False
    data.sheet_view.showGridLines = False
    dash.sheet_view.showGridLines = False

    palette["A1"] = "Paleta STG · tema Office"
    palette["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=PURPLE)
    palette["A2"] = "Colores corporativos y complementarios para tablas, gráficos y formas."
    palette["A2"].font = Font(name="Aptos", size=11, color=GRAPHITE)
    swatches = [
        ("Púrpura STG", PURPLE, "Corporativo / Accent 1"),
        ("Naranja STG", ORANGE, "Corporativo / Accent 2"),
        ("Azul petróleo", PETROL, "Complementario / Accent 3"),
        ("Verde salvia", SAGE, "Complementario / Accent 4"),
        ("Arena", SAND, "Complementario / Accent 5"),
        ("Gris medio", MIDGREY, "Complementario / Accent 6"),
        ("Grafito", GRAPHITE, "Texto"),
        ("Gris claro", LIGHT, "Fondos"),
    ]
    palette.append([])
    palette.append(["Nombre", "Muestra", "HEX", "Uso"])
    for c in palette[4]:
        c.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=PURPLE)
        c.alignment = Alignment(vertical="center")
    for name, hexc, use in swatches:
        palette.append([name, "", f"#{hexc}", use])
        row = palette.max_row
        palette.cell(row, 1).font = Font(name="Aptos", size=11)
        palette.cell(row, 2).fill = PatternFill("solid", fgColor=hexc)
        palette.cell(row, 3).font = Font(name="Aptos", size=10, color=MIDGREY)
        palette.cell(row, 4).font = Font(name="Aptos", size=10)
        palette.row_dimensions[row].height = 25
    palette.column_dimensions["A"].width = 24
    palette.column_dimensions["B"].width = 18
    palette.column_dimensions["C"].width = 14
    palette.column_dimensions["D"].width = 30

    headers = ["Categoría", "2024", "2025", "Variación", "Participación 2025"]
    rows = [
        ["Atención", 62, 82, None, None],
        ["Formación", 54, 64, None, None],
        ["Difusión", 43, 47, None, None],
        ["Estudios", 28, 31, None, None],
        ["Coordinación", 15, 18, None, None],
    ]
    for col, value in enumerate(headers, 1):
        cell = data.cell(1, col, value)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=PURPLE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, values in enumerate(rows, 2):
        for c, value in enumerate(values, 1):
            data.cell(r, c, value)
        data.cell(r, 4, f"=C{r}-B{r}")
        data.cell(r, 5, f"=C{r}/SUM($C$2:$C$6)")
        for c in range(1, 6):
            data.cell(r, c).font = Font(name="Aptos", size=11, color=GRAPHITE)
        data.cell(r, 5).number_format = "0%"
    data["A8"] = "Total"
    data["B8"] = "=SUM(B2:B6)"
    data["C8"] = "=SUM(C2:C6)"
    data["D8"] = "=C8-B8"
    for c in range(1, 5):
        data.cell(8, c).font = Font(name="Aptos", size=11, bold=True, color=PURPLE)
    for col, width in {"A":24,"B":14,"C":14,"D":14,"E":20}.items():
        data.column_dimensions[col].width = width
    thin = Side(style="thin", color="D9D6D8")
    for row in data.iter_rows(min_row=1, max_row=8, min_col=1, max_col=5):
        for cell in row:
            cell.border = Border(bottom=thin)

    dash["A1"] = "Dashboard STG"
    dash["A1"].font = Font(name="Aptos Display", size=22, bold=True, color=PURPLE)
    dash["A2"] = "Ejemplo editable · datos vinculados a la hoja Datos"
    dash["A2"].font = Font(name="Aptos", size=11, color=MIDGREY)
    dash["A4"] = "Indicadores"
    dash["A4"].font = Font(name="Aptos", size=10, bold=True, color=ORANGE)
    dash["A5"] = "Total 2025"
    dash["B5"] = "=Datos!C8"
    dash["D5"] = "Variación total"
    dash["E5"] = "=Datos!D8"
    dash["G5"] = "Categorías"
    dash["H5"] = "=COUNTA(Datos!A2:A6)"
    for cell in ("B5","E5","H5"):
        dash[cell].font = Font(name="Aptos Display", size=24, bold=True, color=PURPLE)

    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Distribución comparada · 2025"
    bar.y_axis.title = ""
    bar.x_axis.title = "Valor"
    bar.height = 7
    bar.width = 13
    cats = Reference(data, min_col=1, min_row=2, max_row=6)
    vals = Reference(data, min_col=3, min_row=1, max_row=6)
    bar.add_data(vals, titles_from_data=True)
    bar.set_categories(cats)
    bar.legend = None
    bar.dLbls = DataLabelList()
    bar.dLbls.showVal = True
    try:
        bar.series[0].graphicalProperties.solidFill = ORANGE
        bar.series[0].graphicalProperties.line.solidFill = ORANGE
    except Exception:
        pass
    dash.add_chart(bar, "A8")

    line = LineChart()
    line.title = "Comparación 2024 / 2025"
    line.height = 7
    line.width = 12
    series = Reference(data, min_col=2, max_col=3, min_row=1, max_row=6)
    line.add_data(series, titles_from_data=True)
    line.set_categories(cats)
    line.legend.position = "b"
    try:
        line.series[0].graphicalProperties.line.solidFill = PURPLE
        line.series[1].graphicalProperties.line.solidFill = PETROL
    except Exception:
        pass
    dash.add_chart(line, "H8")
    for col in range(1, 20):
        dash.column_dimensions[get_column_letter(col)].width = 12

    wb.save(XLSX)
    patch_theme(XLSX, "xl/theme/theme1.xml")

def add_logo(slide, x=0.6, y=0.35, w=2.1):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(x), Inches(y), width=Inches(w))

def add_text(slide, text, x, y, w, h, size=20, color=GRAPHITE, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return box

def add_rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.color.rgb = RGBColor.from_string(line or fill)
    return shp

def build_pptx() -> None:
    PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    # 1. Portada: editable, amplios blancos y bloque corporativo.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_logo(s, 0.65, 0.35, 2.2)
    add_text(s, "TEMA", 0.65, 2.15, 2.0, 0.35, 12, ORANGE, True)
    add_text(s, "Título de la\npresentación", 0.65, 2.55, 7.5, 1.4, 40, PURPLE, False, "Aptos Display")
    add_text(s, "Subtítulo descriptivo o bajada editorial.", 0.68, 4.05, 6.5, 0.6, 21, GRAPHITE)
    add_text(s, "2026", 0.68, 4.88, 2.0, 0.4, 22, ORANGE)
    add_rect(s, 10.35, 4.25, 2.98, 3.25, PURPLE)
    add_rect(s, 11.95, 3.72, 1.38, 0.53, ORANGE)

    # 2. Contenido.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_logo(s)
    add_text(s, "01 · CONTEXTO", 0.65, 1.25, 3.0, 0.3, 12, ORANGE, True)
    add_text(s, "Título de contenido", 0.65, 1.62, 8.8, 0.7, 34, PURPLE, False, "Aptos Display")
    add_text(s, "Texto de apoyo editable para desarrollar antecedentes, objetivos o una idea principal. La tipografía aumenta respecto de la versión anterior y conserva el uso generoso del espacio.", 0.68, 2.65, 7.6, 2.1, 21, GRAPHITE)
    add_rect(s, 9.3, 1.6, 3.1, 3.95, LIGHT, LIGHT)
    add_text(s, "IDEA CLAVE", 9.65, 2.05, 2.3, 0.3, 12, ORANGE, True)
    add_text(s, "Una formulación breve que concentre el criterio o hallazgo principal.", 9.65, 2.55, 2.35, 1.9, 22, PURPLE, True)

    # 3. Separador de sección.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, 13.333333, 7.5, PURPLE)
    add_logo(s, 0.65, 0.35, 2.2)
    add_text(s, "02", 0.7, 2.0, 1.3, 0.75, 36, ORANGE, True, "Aptos Display")
    add_text(s, "Segunda sección", 2.05, 2.0, 7.7, 0.8, 38, WHITE, False, "Aptos Display")
    add_text(s, "Palabra clave · evidencia · análisis · impacto", 2.08, 3.08, 6.6, 0.55, 18, WHITE)

    # 4. Hallazgos con gráfico editable.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_logo(s)
    add_text(s, "03 · HALLAZGOS", 0.65, 1.05, 3.0, 0.3, 12, ORANGE, True)
    add_text(s, "Evidencia y lectura de resultados", 0.65, 1.42, 9.8, 0.65, 32, PURPLE, False, "Aptos Display")
    add_rect(s, 0.65, 2.23, 8.7, 1.05, PURPLE)
    add_text(s, "58%\nPARTICIPACIÓN", 0.95, 2.38, 2.0, 0.7, 24, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "23\nCASOS", 3.65, 2.38, 2.0, 0.7, 24, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "4,2\nÍNDICE", 6.55, 2.38, 2.0, 0.7, 24, ORANGE, True, align=PP_ALIGN.CENTER)
    chart_data = ChartData()
    chart_data.categories = ["Categoría A","Categoría B","Categoría C","Categoría D","Categoría E"]
    chart_data.add_series("Valor", (82,64,47,31,18))
    chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.8), Inches(3.75), Inches(6.3), Inches(2.7), chart_data).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.reverse_order = True
    try:
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(ORANGE)
        chart.series[0].format.line.color.rgb = RGBColor.from_string(ORANGE)
    except Exception:
        pass
    add_text(s, "LECTURA", 8.1, 3.85, 1.8, 0.3, 12, ORANGE, True)
    add_text(s, "Hallazgo principal", 8.1, 4.22, 3.7, 0.45, 22, PURPLE, True)
    add_text(s, "Describa qué muestra la visualización, qué comparación es relevante y cuáles son los límites de la inferencia.", 8.1, 4.78, 3.7, 1.35, 18, GRAPHITE)

    # 5. Cierre.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, 13.333333, 7.5, PURPLE)
    add_logo(s, 0.65, 0.4, 2.2)
    add_text(s, "Cierre", 0.7, 2.35, 5.0, 0.7, 40, WHITE, False, "Aptos Display")
    add_text(s, "secretariadegenero.pjud.cl", 0.72, 5.95, 4.4, 0.4, 18, WHITE)
    add_rect(s, 0.72, 5.55, 0.55, 0.05, ORANGE)

    # Metadata and editable Office objects.
    prs.core_properties.title = "Presentación maestra STGND v9.5"
    prs.core_properties.subject = "Patrón editable con identidad STG"
    prs.core_properties.company = "Poder Judicial de Chile" if hasattr(prs.core_properties, "company") else None
    prs.save(PPTX)
    patch_theme(PPTX, "ppt/theme/theme1.xml")

if __name__ == "__main__":
    build_xlsx()
    build_pptx()
    print(f"Generado: {XLSX}")
    print(f"Generado: {PPTX}")
