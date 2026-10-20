from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from zipfile import ZipFile, ZIP_DEFLATED
import shutil
import tempfile
import json
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'base'/'originales'/'STGND - Oficio_original.zip'
OUTPUT=ROOT/'base'/'STGND - Oficio.zip'
QA_REPORT=ROOT/'base'/'ALFRED_QA_v97.json'
PURPLE='4C2B46'; ORANGE='E97700'; GREY='686868'; LIGHT='F4F1F3'; BLACK='222222'

def set_font(run,name='Aptos',size=11,color=BLACK,bold=None):
    run.font.name=name
    run.font.size=Pt(size)
    run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    rpr=run._r.get_or_add_rPr()
    rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts'); rpr.insert(0,rfonts)
    for attr in ['ascii','hAnsi','eastAsia','cs']:
        rfonts.set(qn('w:'+attr),name)

def p_border(p,side='bottom',color=ORANGE,size=8,space=3):
    pPr=p._p.get_or_add_pPr()
    pBdr=pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr=OxmlElement('w:pBdr'); pPr.append(pBdr)
    old=pBdr.find(qn('w:'+side))
    if old is not None: pBdr.remove(old)
    b=OxmlElement('w:'+side)
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),str(size))
    b.set(qn('w:space'),str(space)); b.set(qn('w:color'),color)
    pBdr.append(b)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill)

def set_cell_border(cell,**edges):
    tcPr=cell._tc.get_or_add_tcPr()
    borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge,opts in edges.items():
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        for k,v in opts.items(): el.set(qn('w:'+k),str(v))

def set_margins(doc,top,bottom,left,right,header=8,footer=8):
    for s in doc.sections:
        s.top_margin=Mm(top); s.bottom_margin=Mm(bottom)
        s.left_margin=Mm(left); s.right_margin=Mm(right)
        s.header_distance=Mm(header); s.footer_distance=Mm(footer)

def style_all_runs(doc):
    for p in doc.paragraphs:
        for r in p.runs: set_font(r,'Aptos',11,BLACK)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs: set_font(r,'Aptos',10.5,BLACK)

def style_heading(p,size=11):
    for r in p.runs: set_font(r,'Aptos',size,PURPLE,True)
    p.paragraph_format.space_before=Pt(8)
    p.paragraph_format.space_after=Pt(4)
    p_border(p,'bottom',ORANGE,6,2)

def style_meta(p,size=10.5):
    for r in p.runs: set_font(r,'Aptos',size,BLACK)
    p.paragraph_format.space_after=Pt(2)

def style_logo_paras(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p._p.xpath('.//w:drawing'):
                        p.paragraph_format.space_after=Pt(3)
                        p_border(p,'bottom',ORANGE,5,2)
                        return

def improve_oficio(doc):
    # Horizontal margins stay at 30 mm because official tab stops depend on them.
    # Only vertical whitespace and typography are tightened to keep one-page output.
    set_margins(doc,18,18,30,30)
    style_all_runs(doc); style_logo_paras(doc)
    ps=doc.paragraphs
    for i,p in enumerate(ps):
        txt=p.text.strip()
        if i<=5:
            for r in p.runs:
                set_font(r,'Aptos',10.5,PURPLE if txt.startswith(('A','DE')) else BLACK,True)
            p.paragraph_format.space_after=Pt(1.5)
        elif 'Sin otro particular' in txt or 'Agradeciendo' in txt:
            for r in p.runs: set_font(r,'Aptos',11,BLACK)
            p.paragraph_format.space_before=Pt(6)
        elif txt.startswith('Distribución') or txt.startswith('- '):
            for r in p.runs: set_font(r,'Aptos',9,GREY)
            p.paragraph_format.space_after=Pt(0)
        else:
            for r in p.runs: set_font(r,'Aptos',11,BLACK)
            p.paragraph_format.line_spacing=1.05
            p.paragraph_format.space_after=Pt(3)
    if doc.tables:
        t=doc.tables[0]
        for ri,row in enumerate(t.rows):
            for ci,cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_font(r,'Aptos',10.5,PURPLE if ci==0 else BLACK,ci==0)
            if ri==0:
                for cell in row.cells:
                    set_cell_border(cell,bottom={'val':'single','sz':'6','color':ORANGE,'space':'0'})

def improve_agenda(doc):
    set_margins(doc,20,20,22,22); style_all_runs(doc)
    ps=doc.paragraphs
    if len(ps)>1:
        for r in ps[1].runs: set_font(r,'Aptos Display',16,PURPLE,False)
        ps[1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        p_border(ps[1],'bottom',ORANGE,7,3)
    for idx in [7,9,14]:
        if idx<len(ps): style_heading(ps[idx],11)
    for i,p in enumerate(ps):
        if i not in [1,7,9,14]: style_meta(p,10.5 if i<7 else 11)

def improve_acta(doc):
    set_margins(doc,18,18,20,20); style_all_runs(doc)
    ps=doc.paragraphs
    if ps:
        for r in ps[0].runs: set_font(r,'Aptos Display',15,PURPLE,False)
        ps[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        p_border(ps[0],'bottom',ORANGE,7,3)
    for p in ps:
        txt=p.text.strip().upper()
        if txt in {'OBJETIVOS DE LA REUNIÓN:','AGENDA DE LA REUNION','DOCUMENTOS PREPARATORIOS','DESARROLLO DE LA REUNIÓN:','ACUERDOS','APROBACIÓN DEL ACTA'}:
            style_heading(p,10.5)
        else:
            style_meta(p,10)
    for t in doc.tables:
        for ri,row in enumerate(t.rows):
            for cell in row.cells:
                if ri==0:
                    set_cell_shading(cell,PURPLE)
                    for p in cell.paragraphs:
                        for r in p.runs: set_font(r,'Aptos',8.5,'FFFFFF',True)
                else:
                    for p in cell.paragraphs:
                        for r in p.runs: set_font(r,'Aptos',9,BLACK)
            if ri==0:
                for cell in row.cells:
                    set_cell_border(cell,bottom={'val':'single','sz':'8','color':ORANGE,'space':'0'})

def improve_informe(doc):
    set_margins(doc,18,18,22,22); style_all_runs(doc)
    for p in doc.paragraphs:
        txt=p.text.strip()
        if txt in {'Antecedentes','Introducción','Análisis','Consideraciones Finales.'}:
            for r in p.runs: set_font(r,'Aptos',11,PURPLE,True)
        elif txt.lower().startswith('contenido'):
            for r in p.runs: set_font(r,'Aptos Display',15,PURPLE,False)
            p_border(p,'bottom',ORANGE,7,3)
        elif txt:
            for r in p.runs: set_font(r,'Aptos',11,BLACK)
            p.paragraph_format.line_spacing=1.08
            p.paragraph_format.space_after=Pt(4)

def improve_minuta(doc):
    set_margins(doc,20,20,22,22); style_all_runs(doc)
    ps=doc.paragraphs
    for i,p in enumerate(ps):
        txt=p.text.strip()
        if i in (0,13,21):
            for r in p.runs: set_font(r,'Aptos Display',15,PURPLE,False)
            p_border(p,'bottom',ORANGE,7,3)
            p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.space_after=Pt(7)
        elif ':' in txt and i<18:
            for r in p.runs: set_font(r,'Aptos',10.5,BLACK)
            p.paragraph_format.space_after=Pt(2)
        elif txt:
            for r in p.runs: set_font(r,'Aptos',11,BLACK)
            p.paragraph_format.line_spacing=1.08
            p.paragraph_format.space_after=Pt(4)

def soften_signature_placeholder(path):
    tmp=path.with_suffix('.tmp.docx')
    with ZipFile(path,'r') as zin, ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename.endswith('.xml'):
                try:
                    text=data.decode('utf-8')
                    # Preserve geometry/object type. Only replace legacy Office colors.
                    text=text.replace('729FCF','F4F1F3').replace('3465A4','D8CFD6').replace('E36C0A','E97700')
                    data=text.encode('utf-8')
                except UnicodeDecodeError:
                    pass
            zout.writestr(item,data)
    tmp.replace(path)


def normalize_header_graphics(path,right_align=False):
    """Remove decorative floating header shapes and convert header pictures to inline."""
    W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
    A='http://schemas.openxmlformats.org/drawingml/2006/main'
    ns={'w':W,'wp':WP,'a':A}
    tmp=path.with_suffix('.header.docx')
    with ZipFile(path,'r') as zin, ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename.startswith('word/header') and item.filename.endswith('.xml'):
                root=etree.fromstring(data)
                changed=False
                for anchor in list(root.xpath('.//wp:anchor',namespaces=ns)):
                    if anchor.xpath('.//a:blip',namespaces=ns):
                        inline=etree.Element('{%s}inline'%WP)
                        for attr in ['distT','distB','distL','distR']:
                            inline.set(attr,'0')
                        for tag in ['extent','effectExtent','docPr','cNvGraphicFramePr']:
                            el=anchor.find('{%s}%s'%(WP,tag))
                            if el is not None:
                                inline.append(el)
                        graphic=anchor.find('{%s}graphic'%A)
                        if graphic is not None:
                            inline.append(graphic)
                        anchor.getparent().replace(anchor,inline)
                        changed=True
                    else:
                        run=anchor
                        while run is not None and run.tag != '{%s}r'%W:
                            run=run.getparent()
                        if run is not None and run.getparent() is not None:
                            run.getparent().remove(run)
                            changed=True
                if right_align:
                    for p in root.xpath('.//w:p[.//wp:inline]',namespaces=ns):
                        pPr=p.find('{%s}pPr'%W)
                        if pPr is None:
                            pPr=etree.Element('{%s}pPr'%W)
                            p.insert(0,pPr)
                        jc=pPr.find('{%s}jc'%W)
                        if jc is None:
                            jc=etree.SubElement(pPr,'{%s}jc'%W)
                        jc.set('{%s}val'%W,'right')
                if changed:
                    data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
            zout.writestr(item,data)
    tmp.replace(path)

def drawing_counts(path):
    WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
    ns={'wp':WP}
    body_anchors=0
    header_anchors=0
    header_inlines=0
    with ZipFile(path) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        body_anchors=len(xml.xpath('.//wp:anchor',namespaces=ns))
        for part in z.namelist():
            if part.startswith('word/header') and part.endswith('.xml'):
                hx=etree.fromstring(z.read(part))
                header_anchors += len(hx.xpath('.//wp:anchor',namespaces=ns))
                header_inlines += len(hx.xpath('.//wp:inline',namespaces=ns))
    return body_anchors,header_anchors,header_inlines

def audit_documents(src,out,specs):
    required={
        'STGND - Oficio.docx':(18,18,30,30),
        'STGND - Oficio circular.docx':(18,18,30,30),
        'STGND - Agenda.docx':(20,20,22,22),
        'STGND - Acta.docx':(18,18,20,20),
        'STGND - Informe.docx':(18,18,22,22),
        'STGND - Minuta.docx':(20,20,22,22),
    }
    rows=[]
    for name,_ in specs:
        original=Document(src/name)
        current=Document(out/name)
        if len(original.paragraphs)!=len(current.paragraphs) or len(original.tables)!=len(current.tables):
            raise RuntimeError(f'Estructura alterada en {name}')
        s=current.sections[0]
        margins=tuple(round(x.mm,1) for x in [s.top_margin,s.bottom_margin,s.left_margin,s.right_margin])
        if margins != required[name]:
            raise RuntimeError(f'Márgenes inesperados en {name}: {margins}')
        body_anchors,header_anchors,header_inlines=drawing_counts(out/name)
        if body_anchors != 1:
            raise RuntimeError(f'{name}: se esperaba solo el ancla técnica FIRMADIGITAL; encontradas {body_anchors}')
        if header_anchors != 0:
            raise RuntimeError(f'{name}: quedan {header_anchors} objetos flotantes en cabecera')
        rows.append({
            'archivo':name,
            'pagina_mm':[round(s.page_width.mm,1),round(s.page_height.mm,1)],
            'margenes_mm':{'superior':margins[0],'inferior':margins[1],'izquierdo':margins[2],'derecho':margins[3]},
            'cabecera_pie_mm':[round(s.header_distance.mm,1),round(s.footer_distance.mm,1)],
            'parrafos':len(current.paragraphs),
            'tablas':len(current.tables),
            'anclas_cuerpo':body_anchors,
            'anclas_cabecera':header_anchors,
            'imagenes_inline_cabecera':header_inlines,
        })
    return {
        'version':'9.7',
        'documentos_activos':len(rows),
        'criterios':{
            'margen_lateral_minimo_general_mm':20,
            'margen_superior_inferior_minimo_general_mm':18,
            'oficios_margen_lateral_mm':30,
            'unica_ancla_flotante_admitida':'FIRMADIGITAL de origen Alfred',
            'graficos_cabecera':'solo imágenes inline; sin formas flotantes',
            'ficha_compras':'excluida del paquete Alfred activo',
        },
        'documentos':rows,
    }

def main():
    if not SOURCE.exists():
        raise SystemExit(f'Missing authoritative source: {SOURCE}')

    with tempfile.TemporaryDirectory() as td:
        td=Path(td); src=td/'src'; out=td/'out'
        src.mkdir(); out.mkdir()
        with ZipFile(SOURCE) as z:
            z.extractall(src)

        specs=[
            ('STGND - Oficio.docx',improve_oficio),
            ('STGND - Oficio circular.docx',improve_oficio),
            ('STGND - Agenda.docx',improve_agenda),
            ('STGND - Acta.docx',improve_acta),
            ('STGND - Informe.docx',improve_informe),
            ('STGND - Minuta.docx',improve_minuta),
        ]
        for name,fn in specs:
            doc=Document(src/name)
            fn(doc)
            doc.save(out/name)
            soften_signature_placeholder(out/name)
            if name=='STGND - Informe.docx':
                normalize_header_graphics(out/name,False)
            elif name=='STGND - Minuta.docx':
                normalize_header_graphics(out/name,True)

        qa=audit_documents(src,out,specs)

        tmpzip=td/'alfred.zip'
        with ZipFile(tmpzip,'w',ZIP_DEFLATED) as z:
            for name,_ in specs:
                z.write(out/name,name)
        shutil.copy2(tmpzip,OUTPUT)
        QA_REPORT.write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
        print(f'Updated {OUTPUT}')
        print(f'QA report: {QA_REPORT}')

if __name__=='__main__':
    main()
